# -*- coding: utf-8 -*-
"""Construit docs/Sijil-Presentation.pptx (présentation commerciale). Usage : python3 docs/build/build_deck.py"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
CAP = lambda n: os.path.join(ROOT, 'docs', 'captures', n)
NAVY = RGBColor(0x0E, 0x43, 0x77); BLUE = RGBColor(0x16, 0x59, 0x8D); TEAL = RGBColor(0x2F, 0xB5, 0xB4); PAPER = RGBColor(0xF4, 0xF1, 0xE8); INK = RGBColor(0x12, 0x26, 0x3A); GREY = RGBColor(0x6B, 0x77, 0x85); WHITE = RGBColor(255, 255, 255)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]
FONT = 'Calibri'

def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); r.shadow.inherit = False
    return r

def text(slide, x, y, w, h, s, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = FONT; r.font.italic = italic
    return tb

def bullets(slide, x, y, w, items, size=18, gap=10, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
        head, _, tail = it.partition(' — ')
        r = p.add_run(); r.text = '● '; r.font.size = Pt(size - 4); r.font.color.rgb = TEAL; r.font.name = FONT
        r = p.add_run(); r.text = head; r.font.size = Pt(size); r.font.bold = bool(tail); r.font.color.rgb = NAVY if tail else color; r.font.name = FONT
        if tail:
            r = p.add_run(); r.text = ' — ' + tail; r.font.size = Pt(size - 1); r.font.color.rgb = color; r.font.name = FONT
    return tb

def pic(slide, name, x, y, w=None, h=None, border=True):
    path = CAP(name) if not os.path.isabs(name) else name
    im = Image.open(path); ar = im.size[1] / im.size[0]
    if w and not h: h = w * ar
    if h and not w: w = h / ar
    if w and h and h / w > ar: h = w * ar
    elif w and h: w = h / ar
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if border: p.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE); p.line.width = Pt(0.75)
    return p

def title(slide, t, sub=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    text(slide, 0.6, 0.35, 12.1, 0.8, t, size=32, color=NAVY, bold=True)
    if sub: text(slide, 0.6, 1.1, 12.1, 0.5, sub, size=17, color=GREY)
    text(slide, 0.6, 7.08, 9, 0.3, '© 2026 Rayane M. — Tous droits réservés', size=10, color=GREY)

def new(bgc=WHITE):
    s = prs.slides.add_slide(BL); bg(s, bgc); return s

def card(slide, x, y, w, h, head, body, fill=WHITE, hc=NAVY, size=15):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); r.adjustments[0] = 0.06; r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.color.rgb = RGBColor(0xD9, 0xDF, 0xE5); r.shadow.inherit = False
    text(slide, x + 0.2, y + 0.15, w - 0.4, 0.5, head, size=17, color=hc, bold=True)
    text(slide, x + 0.2, y + 0.75, w - 0.4, h - 0.9, body, size=size, color=INK)

# 1 couverture
s = new(PAPER)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height); band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
pic(s, os.path.join(ROOT, 'brand', 'logo-sijil-complet.png'), 0.9, 1.15, h=5.1, border=False)
text(s, 5.6, 1.9, 7.2, 1.4, 'Gestion du temps de travail', size=40, color=NAVY, bold=True)
text(s, 5.6, 3.55, 7.0, 1.2, "La mise à disposition et le personnel de prestation de services, de A à Z : du bordereau des prix jusqu'à l'attachement et à la facture.", size=20, color=INK)
text(s, 5.6, 5.4, 7, 0.5, 'En français, en arabe et en anglais · sur ordinateur et sur téléphone', size=16, color=BLUE, bold=True)
text(s, 5.6, 6.6, 7, 0.4, "Version 3.25.1 · Octobre 2026 · Captures : entreprise fictive", size=11, color=GREY)

# 2 problème
s = new(); title(s, "Un personnel en rotation, c'est vite le désordre", "Fichiers Excel, messages, papiers : chacun a sa version")
card(s, 0.6, 1.9, 3.9, 2.4, 'Des jours oubliés', "On pointe quand on peut. Au moment de facturer, il manque des jours et personne ne sait lesquels.")
card(s, 4.7, 1.9, 3.9, 2.4, 'Des calculs à la main', "Soldes de congé, dates de reprise, quantités de l'attachement : chaque calcul est une occasion d'erreur.")
card(s, 8.8, 1.9, 3.9, 2.4, 'Des documents perdus', "Fiches de paie, titres de congé, attestations : dans une boîte mail, un dossier, un téléphone…")
card(s, 0.6, 4.55, 3.9, 2.25, 'Un client dans le flou', "Il demande « qui était présent mardi ? » et attend une réponse par mail.")
card(s, 4.7, 4.55, 3.9, 2.25, 'Des demandes éparpillées', "Chaque agent écrit à son responsable, qui relaie à la direction, sans suivi.")
card(s, 8.8, 4.55, 3.9, 2.25, 'Des prévisions impossibles', "Qui part en congé le mois prochain ? Qui reprend ? Il faut tout recompter.")

# 3 solution
s = new(); title(s, 'Sijil : un seul outil, simple, pour tout le monde', 'Chacun voit ce qui le concerne, rien de plus')
pic(s, '12-global-mois.png', 0.6, 1.75, w=8.3)
bullets(s, 9.2, 1.9, 3.7, ["Un pointage lisible — vert, orange, rouge", "Des prévisions — calculées toutes seules", "Un solde — à jour en temps réel", "Un client informé — sans rien demander"], size=17, gap=14)


# 3b de A à Z
s = new(PAPER); title(s, "De A à Z : du bordereau des prix à la facture", "Sijil gère la mise à disposition et le personnel de prestation de services de bout en bout")
steps = [("1", "Bordereau des prix", "Déposé et lu automatiquement : fonctions, nombres, délais, prix"), ("2", "Contrat et effectifs", "Début, durée, fin ; effectifs bornés par le contrat"), ("3", "Agents et véhicules", "Affectés au contrat, accès envoyés par e-mail"), ("4", "Pointage", "Jour après jour, avec prévisions, congés, soldes et documents"), ("5", "Attachement", "Édité sur les jours pointés, puis validé et figé"), ("6", "Facture", "Établie à partir de l'attachement validé, avec l'historique")]
for i, (n, h_, b) in enumerate(steps):
    x = 0.6 + i * 2.07
    sh = s.shapes.add_shape(MSO_SHAPE.PENTAGON if i < 5 else MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(1.95), Inches(1.0)); sh.fill.solid(); sh.fill.fore_color.rgb = NAVY if i % 2 == 0 else BLUE; sh.line.fill.background(); sh.shadow.inherit = False
    text(s, x + 0.08, 2.1, 1.65, 0.8, n + ". " + h_, size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, 3.2, 1.95, 2.4, b, size=14, color=INK)
text(s, 0.6, 5.7, 12.1, 1.2, ["Un seul outil, une seule saisie : le bordereau des prix lu au départ alimente le contrat, les effectifs, le pointage, l'attachement et la facture.", "Le client suit le pointage et laisse ses remarques ; vos données restent dans votre classeur Google."], size=17, color=NAVY, bold=True)

# 4 profils
s = new(); title(s, 'Quatre profils, quatre espaces', "Chaque personne retrouve l'essentiel dès l'ouverture")
for i, (h_, b) in enumerate([("L'agent", ["Sa situation du jour", "Son solde au départ", "Son pointage sur 1 an", "Ses demandes et documents"]), ('Le responsable d’équipe', ["Pointe son équipe", "Traite les demandes", "Dépose les documents", "Gère ses agents"]), ("L'administrateur", ["Contrats et bordereaux", "Attachements et factures", "Export, import, sauvegarde", "Réglages de l'entreprise"]), ('Le client', ["Consulte le pointage", "Voit contrat et factures", "Laisse des remarques", "Ne modifie rien"])]):
    card(s, 0.6 + i * 3.1, 1.9, 2.95, 4.6, h_, ['● ' + x for x in b], size=16, hc=[BLUE, NAVY, NAVY, TEAL][i])

# 5 pointage
s = new(); title(s, 'Pointer en un clic', "T, R, ABS : sur un jour, sur une période, ou directement dans la grille")
pic(s, '10-pointer.png', 0.6, 1.75, w=8.3)
bullets(s, 9.2, 1.9, 3.7, ["Prévisions — dès le premier pointage, selon la rotation", "Rotation par agent — 28/28, 14/14, 21/21, 6/2, 3/3…", "Un pointage réel — n'est jamais écrasé", "Véhicules — pointés comme des agents"], size=16, gap=12)

# 6 compléter
s = new(PAPER); title(s, "Vous avez oublié de pointer ? Un bouton suffit", "« Copier le dernier pointage » rattrape les jours manquants")
pic(s, '11-completer.png', 0.6, 1.75, w=8.3)
bullets(s, 9.2, 1.9, 3.7, ["Agents cochés — ou tous en une case", "Dernier statut recopié — jusqu'à aujourd'hui", "Résumé avant d'écrire — rien de caché", "Jours déjà pointés — jamais remplacés"], size=16, gap=12)

# 7 vues
s = new(); title(s, "Voir loin : 3 mois, 6 mois, 1 an", "Les départs, les reprises et les fêtes, en un coup d'œil")
pic(s, '14-vue-1-an.png', 0.6, 1.7, w=8.6)
bullets(s, 9.5, 1.9, 3.5, ["Jours fériés algériens — fêtes nationales et religieuses", "Jour courant — repéré, la grille s'y place", "Totaux T, CR, ABS — et solde cumulé", "Dates religieuses — corrigeables en 2 clics"], size=16, gap=12)

# 8 agent
s = new(PAPER); title(s, "L'agent sait toujours où il en est", "Plus besoin d'appeler pour demander son solde ou sa date de reprise")
pic(s, '60-agent-accueil.png', 0.6, 1.75, w=7.4); pic(s, '93-mobile.png', 8.4, 1.6, h=5.2)
text(s, 11.0, 2.3, 2.2, 3.2, ["« Le jour de votre départ, votre solde sera de +43 jours. »"], size=15, color=NAVY, italic=True)

# 9 demandes & documents
s = new(); title(s, 'Demandes et documents, sans courir après', 'Du titre de congé à la fiche de paie')
pic(s, '70-demandes-chef.png', 0.6, 1.75, w=5.9); pic(s, '64-documents-agent.png', 6.9, 1.75, w=5.9)
bullets(s, 0.6, 5.5, 6.1, ["Demandes regroupées par thème — une seule demande part à la direction", "Date de reprise calculée — départ + durée du repos"], size=15, gap=8)
bullets(s, 6.9, 5.5, 6.1, ["Documents rangés tout seuls — type, période, nom, dossier Drive", "Pastille et e-mail — l'agent est prévenu"], size=15, gap=8)

# 10 contrat
s = new(PAPER); title(s, 'Du contrat à la facture, sans ressaisie', 'Le bordereau des prix se lit tout seul, le reste en découle')
pic(s, os.path.join(ROOT, 'docs/captures/crops/40a-contrat.png'), 0.6, 1.7, w=6.6)
bullets(s, 8.1, 1.9, 4.9, ["Bordereau des prix lu automatiquement — PDF scanné, photo, Excel, CSV", "Contrat complet — début, durée, fin visibles de tous", "Effectifs bornés — voyant complet / manque / dépassé", "Historique — par fonction, avec le reste à facturer"], size=16, gap=14)

# 11 attachement
s = new(); title(s, "Attachement et facture, sur les jours pointés", "Validez : le mois est figé. Ensuite, la facture")
pic(s, os.path.join(ROOT, 'docs/captures/crops/50-exports.png'), 0.6, 1.7, w=6.5)
bullets(s, 8.3, 1.9, 4.7, ["Fiche de pointage — Excel et PDF du mois", "Attachement — quantités sur les jours T pointés", "Validation — mêmes quantités, mêmes prix", "Facture — numéro, date, PDF et Excel", "Brouillon du mois suivant — à titre indicatif"], size=16, gap=12)

# 12 client
s = new(PAPER); title(s, 'Votre client suit, et vous répond', 'Un compte de consultation créé tout seul, limité à son contrat')
pic(s, '81-client-remarque.png', 0.6, 1.75, w=5.9); pic(s, '83-remarque-prestataire.png', 6.9, 1.75, w=5.9)
bullets(s, 0.6, 5.5, 6.1, ["Un clic sur un jour — le client laisse sa remarque", "Il ne modifie rien — lecture seule"], size=15, gap=8)
bullets(s, 6.9, 5.5, 6.1, ["Repère rouge dans la grille — et e-mail au responsable d’équipe", "« Vu » et réponse — le client voit la réponse"], size=15, gap=8)

# 13 notifications
s = new(); title(s, 'Tout le monde est prévenu au bon moment', 'Pastilles, messages à l\'écran et e-mails')
card(s, 0.6, 1.9, 3.9, 2.5, 'Nouveau document', "Pastille « Mes documents », message à l'écran et e-mail avec le nom du fichier.")
card(s, 4.7, 1.9, 3.9, 2.5, 'Nouvelle demande', "Le responsable d’équipe est prévenu avec les dates et la date de reprise.")
card(s, 8.8, 1.9, 3.9, 2.5, 'Nouvelle remarque', "Le responsable d’équipe voit la pastille et reçoit la remarque du client.")
card(s, 0.6, 4.65, 6.0, 2.1, 'Accès envoyés par e-mail', "À la création d'un agent, ou à tout moment : lien, identifiant et mot de passe provisoire.")
card(s, 6.8, 4.65, 5.9, 2.1, 'Demandes groupées', "Une seule demande par e-mail à la direction, avec la réponse appliquée à toutes.")

# 14 langues
s = new(PAPER); title(s, 'Français, arabe, anglais — et mode sombre', 'Chacun travaille dans sa langue, de droite à gauche en arabe')
pic(s, '92-arabe.png', 0.6, 1.75, w=5.9); pic(s, '91-mode-sombre.png', 6.9, 1.75, w=5.9)
text(s, 0.6, 5.75, 12.1, 1.5, ["Application installable sur téléphone, en plein écran.", "Langue et mode sombre : un petit sélecteur discret en bas à droite."], size=18, color=NAVY)

# 15 sécurité
s = new(); title(s, 'Vos données restent chez vous', 'Un classeur Google par entreprise, jamais mélangé')
bullets(s, 0.6, 1.9, 6.2, ["Un classeur par entreprise — vous pouvez l'ouvrir à tout moment", "Mots de passe — jamais stockés en clair, connexions limitées", "Droits par profil — chacun ne voit que son périmètre", "Textes saisis — jamais exécutés comme des formules", "Sauvegarde et restauration — en un clic", "Export et import Excel — vos données ne sont pas prisonnières", "Charte de confidentialité — visible sur chaque page"], size=16, gap=10)
pic(s, os.path.join(ROOT, 'docs/captures/crops/43-setup-maintenance.png'), 7.6, 1.75, h=5.1)

# 16 toutes les fonctions
s = new(PAPER); title(s, 'Et ce n\'est pas tout', 'Toutes les fonctions utiles au quotidien')
cols = [["Rattraper les jours oubliés", "Vues 1 mois, 3, 6, 12 mois", "Jours fériés algériens", "Solde T − CR et solde au départ", "Rotation par agent", "Modification de plusieurs agents", "Véhicules mis à disposition", "Aide intégrée en 3 langues"],
        ["Demandes regroupées", "Date de reprise calculée", "Documents rangés et renommés", "Accès par e-mail", "Pastilles et notifications", "Lecture du bordereau des prix (OCR)", "Effectifs bornés par le contrat", "Compte client et remarques"],
        ["Attachement figé puis facture", "Historique et reste à facturer", "Fiche de pointage Excel / PDF", "Export et import Excel", "Sauvegarde et restauration", "Mode sombre", "Installable sur téléphone", "Charte de confidentialité"]]
for i, c in enumerate(cols):
    tb = s.shapes.add_textbox(Inches(0.6 + i * 4.2), Inches(1.8), Inches(4.0), Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    for j, it in enumerate(c):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); p.space_after = Pt(11)
        r = p.add_run(); r.text = '✔  '; r.font.color.rgb = TEAL; r.font.size = Pt(17); r.font.bold = True; r.font.name = FONT
        r = p.add_run(); r.text = it; r.font.size = Pt(17); r.font.color.rgb = NAVY; r.font.name = FONT

# 17 mise en route
s = new(); title(s, 'Prêt en trois étapes', 'Du premier essai au premier pointage')
for i, (h_, b) in enumerate([('1. Demandez votre essai', "Un formulaire sur la page d'accueil. Votre espace est créé et les accès vous sont envoyés par e-mail."), ('2. Déposez votre contrat', "Le bordereau des prix est lu automatiquement. Ajoutez vos agents : leurs accès partent par e-mail."), ('3. Pointez', "Dès le premier jour. Les prévisions, les soldes et les vues se calculent seuls.")]):
    card(s, 0.6 + i * 4.2, 2.0, 3.95, 3.6, h_, b, size=17)
text(s, 0.6, 6.05, 12.1, 0.6, "Besoin de reprendre un historique ? Importez vos fichiers Excel : un mois par onglet.", size=17, color=BLUE, bold=True)

# 18 fin
s = new(NAVY)
pic(s, os.path.join(ROOT, 'brand', 'logo-sijil-symbole.png'), 5.67, 0.9, h=2.0, border=False)
text(s, 0.5, 3.2, 12.3, 1.0, 'Sijil', size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.5, 4.3, 12.3, 0.7, 'Gestion du temps de travail', size=26, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.5, 5.4, 12.3, 0.6, "Demandez votre essai gratuit depuis la page d'accueil.", size=20, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0.5, 6.7, 12.3, 0.4, '© 2026 Rayane M. — Tous droits réservés', size=11, color=RGBColor(0xB8, 0xC8, 0xD8), align=PP_ALIGN.CENTER)

out = os.path.join(ROOT, 'docs', 'Sijil-Presentation.pptx'); prs.save(out); print('ok', out, len(prs.slides._sldIdLst), 'diapositives')
